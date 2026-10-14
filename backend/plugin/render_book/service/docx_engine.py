#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
纯原生通用 Word 题本渲染引擎
职责：
1. 接收 template.docx 模板路径与业务 context，执行 docxtpl 渲染；
2. 清除渲染过程产生的空段落，保留必要分页符与版式元素；
3. 将渲染好的 .docx 导出为高清矢量 .pdf；
4. 抽取封面与第一页高清预览图 .jpg。
"""

import gc
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Any

from docx import Document
from docxtpl import DocxTemplate, InlineImage
import fitz

from backend.common.log import log


def rebind_inline_images(obj: Any, tpl: DocxTemplate, seen: set | None = None) -> None:
    """递归遍历渲染上下文，将所有 InlineImage 对象安全绑定到当前正在渲染的 DocxTemplate 实例"""
    if seen is None:
        seen = set()
    obj_id = id(obj)
    if obj_id in seen:
        return
    seen.add(obj_id)

    if isinstance(obj, InlineImage):
        obj.tpl = tpl
    elif isinstance(obj, dict):
        for v in obj.values():
            rebind_inline_images(v, tpl, seen)
    elif isinstance(obj, (list, tuple, set)):
        for item in obj:
            rebind_inline_images(item, tpl, seen)


class DocxRenderEngine:
    """通用 Word (.docx) 模板渲染引擎"""

    @staticmethod
    def sanitize_text_nodes(doc: Document) -> int:
        """净化文档中所有 <w:t> 文本节点，清除因变量包含特殊符号（如 <r/>、<x/> 等）导致的 OpenXML Schema 破坏"""
        fixed = 0
        try:
            for t in doc.element.xpath('.//w:t'):
                if len(t) > 0:
                    full_text = "".join(t.itertext())
                    for child in list(t):
                        t.remove(child)
                    t.text = full_text
                    fixed += 1
        except Exception as e:
            log.warning(f"[DocxRenderEngine] 净化 <w:t> 文本节点时发生异常: {e}")
        return fixed

    @staticmethod
    def purge_empty_paragraphs(doc: Document) -> int:
        """物理清除因 Jinja2 标签遗留的空段落，保留硬分页符、分节符与图形"""
        removed = 0
        for p in list(doc.paragraphs):
            # 保留图片、图元、分页符、分节符
            if p._p.xpath('.//w:drawing|.//w:pict|.//w:br[@w:type="page"]|./w:pPr/w:sectPr'):
                continue
            if not p.text.strip():
                p._element.getparent().remove(p._element)
                removed += 1

        # 遍历所有表格单元格，清除多余的空段落（每个单元格至少保留一个段落以合规）
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paras = list(cell.paragraphs)
                    if len(paras) <= 1:
                        continue
                    for p in paras:
                        if len(cell.paragraphs) <= 1:
                            break
                        if p._p.xpath('.//w:drawing|.//w:pict|.//w:br[@w:type="page"]|./w:pPr/w:sectPr'):
                            continue
                        if not p.text.strip():
                            p._element.getparent().remove(p._element)
                            removed += 1
        return removed

    @classmethod
    def render_docx(cls, template_file: Path, context: dict[str, Any], output_docx: Path) -> Path:
        """渲染生成 .docx 文件"""
        if not template_file.exists():
            raise FileNotFoundError(f"模板文件不存在: {template_file}")

        output_docx.parent.mkdir(parents=True, exist_ok=True)
        tpl = DocxTemplate(str(template_file))
        rebind_inline_images(context, tpl)
        tpl.render(context)

        # 1. 自动转译文档中的 LaTeX 数学公式为 Word 原生矢量公式 (OMML)
        try:
            from backend.plugin.render_book.service.formula_converter import formula_converter
            formula_cnt = formula_converter.convert_formulas_in_document(tpl.docx)
            if formula_cnt > 0:
                log.info(f"[DocxRenderEngine] 成功转译 LaTeX 数学公式共 {formula_cnt} 处为 Word 原生矢量公式")
        except Exception as e:
            log.warning(f"[DocxRenderEngine] 公式转译过程发生异常，已安全跳过: {e}")

        # 2. 净化所有文本节点，清除因变量包含特殊符号（如 <r/>、<x/> 等）导致的 OpenXML Schema 破坏
        fixed_text_cnt = cls.sanitize_text_nodes(tpl.docx)
        if fixed_text_cnt > 0:
            log.info(f"[DocxRenderEngine] 完成文本节点底层 Schema 净化，修复异常标签共 {fixed_text_cnt} 处")

        # 3. 渲染后物理清除因 Jinja2 控制流标签遗留的虚无空段落
        removed_cnt = cls.purge_empty_paragraphs(tpl.docx)
        log.info(f"[DocxRenderEngine] 渲染后完成空段落物理净化，剔除无用空行共 {removed_cnt} 处")

        try:
            tpl.save(str(output_docx))
        except PermissionError:
            import time
            output_docx = output_docx.parent / f"{output_docx.stem}_{int(time.time())}.docx"
            tpl.save(str(output_docx))
            log.warning(f"[DocxRenderEngine] 目标 Word 文件被占用，安全另存为: {output_docx}")

        log.info(f"[DocxRenderEngine] 成功渲染 Word 题本: {output_docx}")
        return output_docx

    @staticmethod
    def _find_libreoffice_binary() -> str | None:
        """寻找系统中的 LibreOffice / soffice 可执行文件路径（跨平台，支持 Linux/macOS/Windows）"""
        for cmd in ["libreoffice", "soffice"]:
            found = shutil.which(cmd)
            if found:
                return found

        candidate_paths = [
            "/usr/bin/libreoffice",
            "/usr/bin/soffice",
            "/usr/local/bin/libreoffice",
            "/usr/local/bin/soffice",
            "/opt/libreoffice/program/soffice",
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        for p in candidate_paths:
            p_obj = Path(p)
            if p_obj.is_file():
                if sys.platform != "win32" and not os.access(p, os.X_OK):
                    continue
                return p
        return None

    @classmethod
    def _convert_via_libreoffice(cls, docx_file: Path, output_pdf: Path) -> Path:
        """使用无头 LibreOffice 执行 Word 转 PDF（Linux 生产环境标准）"""
        binary = cls._find_libreoffice_binary()
        if not binary:
            raise RuntimeError(
                "未检测到 LibreOffice 环境 (未找到 libreoffice 或 soffice 命令)。"
                "在 Linux/Docker 环境下请运行: apt-get install -y libreoffice-writer 或 libreoffice"
            )

        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp_dir:
            cmd = [
                binary,
                "--headless",
                "--convert-to",
                "pdf:writer_pdf_Export",
                str(docx_file.resolve()),
                "--outdir",
                str(Path(tmp_dir).resolve()),
            ]
            log.info(f"[DocxRenderEngine] 执行 LibreOffice 转换: {' '.join(cmd)}")
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180)
            if res.returncode != 0:
                err_msg = res.stderr.decode("utf-8", errors="ignore").strip()
                raise RuntimeError(f"LibreOffice 转换异常 (退出码 {res.returncode}): {err_msg}")

            generated_pdf = Path(tmp_dir) / f"{docx_file.stem}.pdf"
            if not generated_pdf.exists() or generated_pdf.stat().st_size == 0:
                raise RuntimeError(f"LibreOffice 未能成功生成中间 PDF: {generated_pdf}")

            try:
                shutil.copyfile(generated_pdf, output_pdf)
            except PermissionError:
                output_pdf = output_pdf.parent / f"{output_pdf.stem}_{int(time.time())}.pdf"
                shutil.copyfile(generated_pdf, output_pdf)
                log.warning(f"[DocxRenderEngine] 目标 PDF 被占用，安全另存为: {output_pdf}")

        if not output_pdf.exists() or output_pdf.stat().st_size == 0:
            raise RuntimeError(f"矢量 PDF 输出校验失败，目标文件不存在: {output_pdf}")

        log.info(f"[DocxRenderEngine] 成功通过 LibreOffice 生成矢量 PDF: {output_pdf}")
        return output_pdf

    @staticmethod
    def _convert_via_msoffice(docx_file: Path, output_pdf: Path) -> Path:
        """使用 Windows 微软 Office COM 接口执行 Word 转 PDF（修复句柄占用与 WinError 32）"""
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        try:
            # ignore_cleanup_errors=True 彻底规避 Windows 下 Word 异步退出时锁住 safe_copy.docx 导致的 [WinError 32]
            with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp_dir:
                temp_docx = Path(tmp_dir) / "safe_copy.docx"
                temp_pdf = Path(tmp_dir) / "safe_copy.pdf"
                shutil.copyfile(docx_file, temp_docx)

                word = None
                doc = None
                try:
                    word = win32com.client.DispatchEx("Word.Application")
                    word.Visible = False
                    word.DisplayAlerts = 0
                    doc = word.Documents.Open(str(temp_docx.resolve()), ReadOnly=1)
                    doc.SaveAs(str(temp_pdf.resolve()), FileFormat=17)  # 17 = wdFormatPDF
                finally:
                    if doc is not None:
                        try:
                            doc.Close(0)
                        except Exception:
                            pass
                        doc = None
                    if word is not None:
                        try:
                            word.Quit()
                        except Exception:
                            pass
                        word = None
                    # 显式解除 Python 侧的 COM RCW 句柄并回收
                    gc.collect()

                if not temp_pdf.exists() or temp_pdf.stat().st_size == 0:
                    raise RuntimeError(f"Word 转 PDF 失败，未生成中间文件: {temp_pdf}")

                try:
                    shutil.copyfile(temp_pdf, output_pdf)
                except PermissionError:
                    output_pdf = output_pdf.parent / f"{output_pdf.stem}_{int(time.time())}.pdf"
                    shutil.copyfile(temp_pdf, output_pdf)
                    log.warning(f"[DocxRenderEngine] 目标 PDF 被占用，安全另存为: {output_pdf}")

            if not output_pdf.exists() or output_pdf.stat().st_size == 0:
                raise RuntimeError(f"矢量 PDF 输出校验失败，目标文件不存在: {output_pdf}")

            log.info(f"[DocxRenderEngine] 成功通过 MS Word 生成矢量 PDF: {output_pdf}")
            return output_pdf
        finally:
            pythoncom.CoUninitialize()

    @classmethod
    def convert_to_pdf(cls, docx_file: Path, output_pdf: Path) -> Path:
        """
        将 .docx 安全转换为高清矢量 .pdf
        跨平台自适应路由：
        1. 若为 Windows 且已安装 Office Word，优先通过 MS Word 转换；
        2. 若在 Linux / Docker / macOS 或未装 Office 的环境中，自动调用 LibreOffice 无头转换。
        """
        output_pdf.parent.mkdir(parents=True, exist_ok=True)

        if sys.platform == "win32":
            try:
                return cls._convert_via_msoffice(docx_file=docx_file, output_pdf=output_pdf)
            except Exception as e:
                if cls._find_libreoffice_binary():
                    log.warning(f"[DocxRenderEngine] 微软 Word COM 转换异常 ({e})，尝试回退至 LibreOffice...")
                    return cls._convert_via_libreoffice(docx_file=docx_file, output_pdf=output_pdf)
                log.error(f"[DocxRenderEngine] 微软 Word COM 转换失败: {e}")
                raise
        else:
            return cls._convert_via_libreoffice(docx_file=docx_file, output_pdf=output_pdf)

    @staticmethod
    def extract_preview_images(pdf_file: Path, output_dir: Path, max_pages: int = 2) -> list[Path]:
        """从生成的 PDF 中抽取前 N 页作为前端展示的缩略图 (.jpg)"""
        output_dir.mkdir(parents=True, exist_ok=True)
        preview_paths: list[Path] = []
        if not pdf_file.exists():
            return preview_paths

        doc = fitz.open(pdf_file)
        pages_to_render = min(len(doc), max_pages)
        for i in range(pages_to_render):
            page = doc[i]
            pix = page.get_pixmap(dpi=150)
            img_path = output_dir / f"preview_page_{i + 1}.jpg"
            pix.save(str(img_path))
            preview_paths.append(img_path)

        doc.close()
        log.info(f"[DocxRenderEngine] 成功抽取预览图共 {len(preview_paths)} 张")
        return preview_paths

    @classmethod
    def render_all(
        cls,
        template_file: Path,
        context: dict[str, Any],
        output_dir: Path,
        base_name: str = "exam_paper",
        compile_pdf: bool = True,
        max_previews: int = 2,
    ) -> dict[str, Any]:
        """
        完整执行渲染流程，返回生成的文件字典
        :return: {'docx': Path, 'pdf': Path | None, 'previews': list[Path]}
        """
        output_dir.mkdir(parents=True, exist_ok=True)
        docx_path = output_dir / f"{base_name}.docx"
        docx_path = cls.render_docx(template_file=template_file, context=context, output_docx=docx_path)

        pdf_path = None
        preview_paths = []
        if compile_pdf:
            pdf_path = output_dir / f"{base_name}.pdf"
            actual_pdf = cls.convert_to_pdf(docx_file=docx_path, output_pdf=pdf_path)
            pdf_path = actual_pdf
            preview_paths = cls.extract_preview_images(pdf_file=pdf_path, output_dir=output_dir, max_pages=max_previews)

        return {
            "docx": docx_path,
            "pdf": pdf_path,
            "previews": preview_paths,
        }


docx_render_engine = DocxRenderEngine()
