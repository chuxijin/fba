# -*- coding: utf-8 -*-
"""
LaTeX 数学公式转 Word 原生矢量公式 (OMML) 转换引擎
采用纯 Python 轻量级解析库 latex2mathml + 微软官方 MML2OMML.XSL 转换器，
自动扫描 Word 文档中的 $...$ (行内公式) 与 $$...$$ (独立公式)，
无缝将其替换为 Word 原生可编辑、高清矢量的 <m:oMath> 节点。
"""

import re
import os
from pathlib import Path
from typing import Any

from lxml import etree
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from loguru import logger as log

try:
    from latex2mathml.converter import convert as latex2mathml_convert
except ImportError:
    latex2mathml_convert = None


# 常见 MML2OMML.XSL 系统路径候选
CANDIDATE_XSL_PATHS = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office15\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office15\MML2OMML.XSL",
]


class FormulaConverter:
    """LaTeX to Word OMML 原生矢量公式转换器"""

    def __init__(self):
        self._xslt_transform = None
        self._init_xslt()

    def _init_xslt(self):
        """加载并预编译 MML2OMML.XSL 转换器"""
        for p in CANDIDATE_XSL_PATHS:
            if os.path.exists(p):
                try:
                    xslt_doc = etree.parse(p)
                    self._xslt_transform = etree.XSLT(xslt_doc)
                    log.info(f"[FormulaConverter] 成功加载微软 MML2OMML.XSL 转换器: {p}")
                    return
                except Exception as e:
                    log.warning(f"[FormulaConverter] 加载 XSLT 失败: {p}, 错误: {e}")

        log.warning("[FormulaConverter] 未在系统中检测到 MML2OMML.XSL，公式转译功能将以降级文本模式运行。")

    @property
    def is_available(self) -> bool:
        return latex2mathml_convert is not None and self._xslt_transform is not None

    def latex_to_omml_element(self, latex_str: str) -> etree.Element | None:
        """
        将单段 LaTeX 公式代码转换为 Word OMML <m:oMath> 节点
        """
        if not self.is_available:
            return None

        # 1. 基础语法清洗与容错
        code = latex_str.strip()
        if code.startswith("$$") and code.endswith("$$"):
            code = code[2:-2].strip()
        elif code.startswith("$") and code.endswith("$"):
            code = code[1:-1].strip()

        if not code:
            return None

        # 中文括号标准化为英文括号
        code = code.replace("（", "(").replace("）", ")")
        # 清理多余斜杠
        code = re.sub(r'\\{2,}', r'\\', code)

        # 2. 转换为 MathML
        try:
            mathml = latex2mathml_convert(code)
        except Exception as e:
            # 常见容错：如果包含未转义的百分号或特殊字符，进行二次尝试
            try:
                fixed_code = re.sub(r'(?<!\\)%', r'\%', code)
                mathml = latex2mathml_convert(fixed_code)
            except Exception:
                log.debug(f"[FormulaConverter] latex2mathml 转换失败: {code}")
                return None

        # 3. MathML 转 OMML
        try:
            mml_doc = etree.fromstring(mathml)
            omml_doc = self._xslt_transform(mml_doc)
            return omml_doc.getroot()
        except Exception as e:
            log.debug(f"[FormulaConverter] XSLT 转换 OMML 失败: {code}, 错误: {e}")
            return None

    def convert_formulas_in_paragraph(self, p: Any) -> int:
        """
        扫描单个段落，若包含 $...$ 或 $$...$$ 公式标记，精准在原地拆解 Run 并嵌入原生公式
        :return: 成功转换的公式数量
        """
        if not self.is_available:
            return 0

        # 如果段落完全不包含 $，零耗时直接跳过
        if "$" not in p.text:
            return 0

        converted_cnt = 0
        runs = list(p.runs)

        for r in runs:
            text = r.text
            if not text or "$" not in text:
                continue

            parts = re.split(r"(\$\$[^\$]+\$\$|\$[^\$]+\$)", text)
            if len(parts) <= 1:
                continue

            rPr_elem = r._r.find(qn("w:rPr"))
            orig_rPr_xml = etree.tostring(rPr_elem, encoding="utf-8").decode("utf-8") if rPr_elem is not None else None

            try:
                insert_idx = p._p.index(r._r)
            except ValueError:
                continue

            p._p.remove(r._r)

            for part in parts:
                if not part:
                    continue

                is_formula = (part.startswith("$$") and part.endswith("$$")) or (part.startswith("$") and part.endswith("$"))
                if is_formula:
                    omml = self.latex_to_omml_element(part)
                    if omml is not None:
                        p._p.insert(insert_idx, etree.fromstring(etree.tostring(omml)))
                        insert_idx += 1
                        converted_cnt += 1
                        continue

                # 非公式纯文字片段（或转换失败降级显示纯文本）
                new_r = docx.text.run.Run(parse_xml(f'<w:r {nsdecls("w")}/>'), p)
                if orig_rPr_xml:
                    new_r._r.append(parse_xml(orig_rPr_xml))
                new_r.text = part
                p._p.insert(insert_idx, new_r._r)
                insert_idx += 1

        return converted_cnt

    def convert_formulas_in_document(self, doc: Any) -> int:
        """
        遍历整篇文档正文与所有表格，自动完成公式转换
        :return: 全书转换成功的公式总数
        """
        if not self.is_available:
            return 0

        total_converted = 0

        # 1. 正文段落
        for p in doc.paragraphs:
            total_converted += self.convert_formulas_in_paragraph(p)

        # 2. 表格单元格段落
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        total_converted += self.convert_formulas_in_paragraph(p)

        return total_converted


formula_converter = FormulaConverter()
