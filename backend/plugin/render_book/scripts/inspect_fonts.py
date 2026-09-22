# -*- coding: utf-8 -*-
import docx

doc = docx.Document(r"D:\100_Work\101_Program\Proj\fba\backend\plugin\render_book\templates\basic_calculation\1.0.0\template.docx")

print("=== Table 1 (Calc row) Fonts ===")
tbl1 = doc.tables[1]
for r in tbl1.rows:
    for c in r.cells:
        for p in c.paragraphs:
            for run in p.runs:
                rPr = run._r.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
                rFonts_dict = {}
                if rPr is not None:
                    rFonts = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                    if rFonts is not None:
                        rFonts_dict = dict(rFonts.attrib)
                print(f"T1 run: text='{run.text}', font={run.font.name}, size={run.font.size}, rFonts={rFonts_dict}")

print("\n=== Table 0 (Punch bar) Fonts ===")
tbl0 = doc.tables[0]
for c in tbl0.rows[0].cells:
    for p in c.paragraphs:
        for run in p.runs:
            rPr = run._r.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
            rFonts_dict = {}
            if rPr is not None:
                rFonts = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                if rFonts is not None:
                    rFonts_dict = dict(rFonts.attrib)
            print(f"T0 run: text='{run.text}', rFonts={rFonts_dict}")
