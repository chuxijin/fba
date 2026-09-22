# -*- coding: utf-8 -*-
"""
公考错题本专属 Word (.docx) 模板生成脚本
定位：公考（行测）错题复盘、思维建模与查漏补缺专属题本
核心特性：
1. 规范单栏 B5 排版，官方仿宋 12pt 与严整悬挂缩进；
2. 每道错题专属结构化【错题深度复盘卡】：
   - (a) 上方左右双栏：左列【错误思路】 vs 右列【正确思路】
   - (b) 下方作答核对栏：我的作答与正确答案
3. 灵活解析呈现模式控制：
   - inline: 错题卡下方紧随官方深度解析
   - appendix: 集中在书末附录呈现，包含答案速查表与逐题解析
4. 原生动态页脚与公式原生矢量支持。
"""

from pathlib import Path
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn


def create_mistake_template():
    target_dir = Path(r"backend/plugin/render_book/templates/gongkao_mistake/1.0.0")
    target_dir.mkdir(parents=True, exist_ok=True)
    docx_path = target_dir / "template.docx"

    doc = docx.Document()

    # 辅助函数：设置东亚字体与英文字体
    def set_font(run, east_asia="仿宋", ascii_font="Times New Roman", size_pt=10.5, bold=False, color_rgb=None):
        run.font.name = ascii_font
        if size_pt:
            run.font.size = Pt(size_pt)
        run.font.bold = bold
        if color_rgb:
            run.font.color.rgb = color_rgb
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{ascii_font}" w:eastAsia="{east_asia}" w:hAnsi="{ascii_font}" w:cs="{east_asia}"/>')
            rPr.append(rFonts)
        else:
            rFonts.set(qn("w:eastAsia"), east_asia)
            rFonts.set(qn("w:ascii"), ascii_font)
            rFonts.set(qn("w:hAnsi"), ascii_font)
            rFonts.set(qn("w:cs"), east_asia)

    # 辅助函数：设置单元格边框
    def set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="C5CAD0", sz="4"):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

    def set_table_borders_none(table):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="none"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:bottom w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="none"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def set_table_left_indent(table, indent_pt: float):
        tblPr = table._tbl.tblPr
        tblInd = tblPr.find(qn('w:tblInd'))
        if tblInd is not None:
            tblPr.remove(tblInd)
        new_tblInd = parse_xml(f'<w:tblInd {nsdecls("w")} w:w="{int(indent_pt * 20)}" w:type="dxa"/>')
        tblPr.append(new_tblInd)

    def set_cell_margins(cell, top=2, bottom=2, left=4, right=4):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>\n'
            f'  <w:top w:w="{int(top * 20)}" w:type="dxa"/>\n'
            f'  <w:bottom w:w="{int(bottom * 20)}" w:type="dxa"/>\n'
            f'  <w:left w:w="{int(left * 20)}" w:type="dxa"/>\n'
            f'  <w:right w:w="{int(right * 20)}" w:type="dxa"/>\n'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_cell_shading(cell, color="F8FAFC"):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
        tcPr.append(shd)

    # 页面版心尺寸（标准 B5: 176mm x 250mm, 页边距左右 20mm, 版心宽 136mm）
    PAGE_WIDTH = Mm(176)
    PAGE_HEIGHT = Mm(250)
    MARGIN_LR = Mm(20)
    CONTENT_WIDTH = PAGE_WIDTH - MARGIN_LR * 2
    HANG_INDENT_PT = 24.0
    OPT_TABLE_WIDTH = CONTENT_WIDTH - Pt(HANG_INDENT_PT)

    # ==================== 第 1 节：封面（公考错题本） ====================
    sec1 = doc.sections[0]
    sec1.page_width = PAGE_WIDTH
    sec1.page_height = PAGE_HEIGHT
    sec1.top_margin = Mm(25)
    sec1.bottom_margin = Mm(20)
    sec1.left_margin = MARGIN_LR
    sec1.right_margin = MARGIN_LR

    # 顶部标签
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badge.paragraph_format.space_before = Pt(20)
    p_badge.paragraph_format.space_after = Pt(12)
    r_badge = p_badge.add_run("★ 错 题 复 盘 · 精 进 提 分 ★")
    set_font(r_badge, east_asia="黑体", size_pt=11.0, bold=True, color_rgb=RGBColor(0x88, 0x11, 0x11))

    # 主标题：公考错题本
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("{{ book.title if book.title else '公考错题本' }}")
    set_font(r_title, east_asia="黑体", size_pt=24, bold=True, color_rgb=RGBColor(0x1A, 0x1A, 0x1A))

    # 副标题
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(40)
    r_sub = p_sub.add_run("{{ book.subtitle if book.subtitle else '错因深度剖析 · 思维正误建模 · 知识盲区击破' }}")
    set_font(r_sub, east_asia="仿宋", size_pt=11.5, color_rgb=RGBColor(0x55, 0x55, 0x55))

    # 错题本信息栏表格
    info_table = doc.add_table(rows=3, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_widths = [CONTENT_WIDTH * 0.5, CONTENT_WIDTH * 0.5]
    info_data = [
        ("学 员 姓 名：{{ metadata.nickname if metadata.nickname else '公务员考生' }}", "收 录 错 题：共 {{ paper.question_count if paper.question_count else 0 }} 道"),
        ("所 属 模 块：{{ metadata.module_title if metadata.module_title else '行测高频错题集' }}", "所 属 题 库：{{ metadata.bank_name if metadata.bank_name else '国家公务员考试真题' }}"),
        ("复 盘 规 划：逐题诊断 · 彻底消化", "打 卡 状 态：[  ] 已掌握    [  ] 需重刷"),
    ]

    for r_idx, row in enumerate(info_table.rows):
        t0, t1 = info_data[r_idx]
        for c_idx, cell in enumerate(row.cells):
            cell.width = info_widths[c_idx]
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="D1D5DB", sz="5")
            set_cell_shading(cell, color="F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.left_indent = Pt(12)
            text = t0 if c_idx == 0 else t1
            r = p.add_run(text)
            set_font(r, east_asia="仿宋", size_pt=10.5, color_rgb=RGBColor(0x33, 0x33, 0x33))

    # 底部名言警句
    p_motto = doc.add_paragraph()
    p_motto.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_motto.paragraph_format.space_before = Pt(100)
    p_motto.paragraph_format.space_after = Pt(12)
    r_motto = p_motto.add_run("“ 错题是最好的老师，复盘是最快的精进。”")
    set_font(r_motto, east_asia="楷体", size_pt=11.5, color_rgb=RGBColor(0x66, 0x66, 0x66))

    # ==================== 第 2 节：错题正文（题目 + 左右双栏复盘卡 + 底部作答栏） ====================
    sec2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    sec2.page_width = PAGE_WIDTH
    sec2.page_height = PAGE_HEIGHT
    sec2.top_margin = Mm(20)
    sec2.bottom_margin = Mm(20)
    sec2.left_margin = MARGIN_LR
    sec2.right_margin = MARGIN_LR

    # 单栏
    sectPr2 = sec2._sectPr
    cols_xml2 = parse_xml(f'<w:cols {nsdecls("w")} w:num="1"/>')
    sectPr2.append(cols_xml2)
    sectPr2.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>'))

    # 正文原生动态页脚
    sec2.footer.is_linked_to_previous = False
    fp2 = sec2.footer.paragraphs[0]
    fp2.text = ""
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp2.paragraph_format.space_before = Pt(6)
    fp2.paragraph_format.space_after = Pt(0)
    pPr2 = fp2._p.get_or_add_pPr()
    pPr2.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))

    r1 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t xml:space="preserve">第 </w:t></w:r>')
    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
    r2 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t xml:space="preserve"> 页  共 </w:t></w:r>')
    fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
    r3 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t> 页</w:t></w:r>')
    fp2._p.append(r1)
    fp2._p.append(fld1)
    fp2._p.append(r2)
    fp2._p.append(fld2)
    fp2._p.append(r3)

    # 题目循环开始
    p_q_loop = doc.add_paragraph()
    p_q_loop.paragraph_format.space_before = Pt(0)
    p_q_loop.paragraph_format.space_after = Pt(0)
    p_q_loop.add_run("{% for q in paper.all_questions %}")

    # 1. 材料引导语段落
    p_mat_intro = doc.add_paragraph()
    p_mat_intro.paragraph_format.space_before = Pt(6)
    p_mat_intro.paragraph_format.space_after = Pt(2)
    p_mat_intro.paragraph_format.left_indent = Pt(HANG_INDENT_PT)
    p_mat_intro.paragraph_format.line_spacing = 1.2
    p_mat_intro.paragraph_format.keep_with_next = True
    p_mat_intro.add_run("{% if q.material and q.material.intro %}")
    r_mi = p_mat_intro.add_run("{{ q.material.intro }}")
    set_font(r_mi, east_asia="黑体", size_pt=10.5, bold=False)
    p_mat_intro.add_run("{% endif %}")

    # 2. 材料图文段落
    doc.add_paragraph("{% if q.material and q.material.blocks %}{% for blk in q.material.blocks %}")
    doc.add_paragraph("{% if not blk.is_img %}")
    p_mat_text = doc.add_paragraph()
    p_mat_text.paragraph_format.space_before = Pt(1)
    p_mat_text.paragraph_format.space_after = Pt(2)
    p_mat_text.paragraph_format.left_indent = Pt(0)
    p_mat_text.paragraph_format.first_line_indent = Pt(HANG_INDENT_PT)
    p_mat_text.paragraph_format.line_spacing = 1.25
    p_mat_text.paragraph_format.keep_with_next = True
    r_mat_txt = p_mat_text.add_run("{{ blk.text }}")
    set_font(r_mat_txt, east_asia="仿宋", size_pt=12.0)
    doc.add_paragraph("{% else %}")
    p_mat_img = doc.add_paragraph()
    p_mat_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mat_img.paragraph_format.space_before = Pt(4)
    p_mat_img.paragraph_format.space_after = Pt(4)
    p_mat_img.paragraph_format.keep_with_next = True
    p_mat_img.add_run("{{ blk.img }}")
    doc.add_paragraph("{% endif %}{% endfor %}{% elif q.material and q.material.lines %}{% for line in q.material.lines %}")
    p_mat_fb = doc.add_paragraph()
    p_mat_fb.paragraph_format.space_before = Pt(1)
    p_mat_fb.paragraph_format.space_after = Pt(2)
    p_mat_fb.paragraph_format.first_line_indent = Pt(HANG_INDENT_PT)
    p_mat_fb.paragraph_format.line_spacing = 1.25
    p_mat_fb.paragraph_format.keep_with_next = True
    r_mat_fb = p_mat_fb.add_run("{{ line }}")
    set_font(r_mat_fb, east_asia="仿宋", size_pt=12.0)
    doc.add_paragraph("{% endfor %}{% endif %}")

    # 3. 题干：悬挂缩进 24pt
    p_stem = doc.add_paragraph()
    p_stem.paragraph_format.space_before = Pt(5)
    p_stem.paragraph_format.space_after = Pt(2)
    p_stem.paragraph_format.line_spacing = 1.25
    p_stem.paragraph_format.keep_with_next = True
    p_stem.paragraph_format.left_indent = Pt(HANG_INDENT_PT)
    p_stem.paragraph_format.first_line_indent = -Pt(HANG_INDENT_PT)
    p_stem.paragraph_format.tab_stops.add_tab_stop(Pt(HANG_INDENT_PT))

    r_no = p_stem.add_run("{{ q.number }}.\t")
    set_font(r_no, east_asia="仿宋", ascii_font="Times New Roman", size_pt=12.0)

    r_tag = p_stem.add_run("{% if q.source_tag %}【{{ q.source_tag }}】{% endif %}")
    set_font(r_tag, east_asia="仿宋", ascii_font="Times New Roman", size_pt=12.0)

    r_stem_0 = p_stem.add_run("{% for seg in q.stem_segments %}{% if seg.bold %}")
    set_font(r_stem_0, east_asia="仿宋", size_pt=12.0)
    r_stem_bold = p_stem.add_run("{{ seg.text }}")
    set_font(r_stem_bold, east_asia="仿宋", size_pt=12.0, bold=True)
    r_stem_else = p_stem.add_run("{% else %}")
    set_font(r_stem_else, east_asia="仿宋", size_pt=12.0)
    r_stem_normal = p_stem.add_run("{{ seg.text }}")
    set_font(r_stem_normal, east_asia="仿宋", size_pt=12.0, bold=False)
    r_stem_end = p_stem.add_run("{% endif %}{% endfor %}")
    set_font(r_stem_end, east_asia="仿宋", size_pt=12.0)

    # 题干配图
    doc.add_paragraph("{% if q.stem_images %}{% for simg in q.stem_images %}")
    p_simg = doc.add_paragraph()
    p_simg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_simg.paragraph_format.space_before = Pt(3)
    p_simg.paragraph_format.space_after = Pt(4)
    p_simg.paragraph_format.keep_with_next = True
    p_simg.add_run("{{ simg }}")
    doc.add_paragraph("{% endfor %}{% endif %}")

    # 4. 4 列选项表格
    doc.add_paragraph("{% if q.opt_cols == 4 %}")
    t4 = doc.add_table(rows=1, cols=4)
    t4.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders_none(t4)
    set_table_left_indent(t4, HANG_INDENT_PT)
    t4.autofit = False
    col_width_4 = OPT_TABLE_WIDTH / 4
    for idx, c in enumerate(t4.rows[0].cells):
        c.width = col_width_4
        set_cell_margins(c, top=2, bottom=2, left=0, right=6)
        cp = c.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.15
        cp.add_run(f"{{% if q.opt_{idx}_img %}}")
        c_pre = cp.add_run(f"{{{{ q.opt_{idx}_prefix }}}} ")
        set_font(c_pre, east_asia="仿宋", size_pt=12.0)
        cp.add_run(f"{{{{ q.opt_{idx}_img }}}}")
        cp.add_run(f"{{% else %}}")
        cr = cp.add_run(f"{{{{ q.opt_{idx} }}}}")
        set_font(cr, east_asia="仿宋", size_pt=12.0)
        cp.add_run(f"{{% endif %}}")
    doc.add_paragraph("{% endif %}")

    # 5. 2 列选项表格
    doc.add_paragraph("{% if q.opt_cols == 2 %}")
    t2 = doc.add_table(rows=2, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders_none(t2)
    set_table_left_indent(t2, HANG_INDENT_PT)
    t2.autofit = False
    col_width_2 = OPT_TABLE_WIDTH / 2
    t2_cells = [t2.rows[0].cells[0], t2.rows[0].cells[1], t2.rows[1].cells[0], t2.rows[1].cells[1]]
    for idx, c in enumerate(t2_cells):
        c.width = col_width_2
        set_cell_margins(c, top=2, bottom=2, left=0, right=6)
        cp = c.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.15
        cp.add_run(f"{{% if q.opt_{idx}_img %}}")
        c_pre = cp.add_run(f"{{{{ q.opt_{idx}_prefix }}}} ")
        set_font(c_pre, east_asia="仿宋", size_pt=12.0)
        cp.add_run(f"{{{{ q.opt_{idx}_img }}}}")
        cp.add_run(f"{{% else %}}")
        cr = cp.add_run(f"{{{{ q.opt_{idx} }}}}")
        set_font(cr, east_asia="仿宋", size_pt=12.0)
        cp.add_run(f"{{% endif %}}")
    doc.add_paragraph("{% endif %}")

    # 6. 1 列选项
    doc.add_paragraph("{% if q.opt_cols == 1 %}")
    for i in range(4):
        doc.add_paragraph(f"{{% if q.opt_{i} or q.opt_{i}_img %}}")
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.space_before = Pt(1)
        p_opt.paragraph_format.space_after = Pt(1)
        p_opt.paragraph_format.left_indent = Pt(HANG_INDENT_PT)
        p_opt.paragraph_format.line_spacing = 1.15
        p_opt.add_run(f"{{% if q.opt_{i}_img %}}")
        p_pre = p_opt.add_run(f"{{{{ q.opt_{i}_prefix }}}} ")
        set_font(p_pre, east_asia="仿宋", size_pt=12.0)
        p_opt.add_run(f"{{{{ q.opt_{i}_img }}}}")
        p_opt.add_run(f"{{% else %}}")
        r_opt = p_opt.add_run(f"{{{{ q.opt_{i} }}}}")
        set_font(r_opt, east_asia="仿宋", size_pt=12.0)
        p_opt.add_run(f"{{% endif %}}")
        doc.add_paragraph("{% endif %}")
    doc.add_paragraph("{% endif %}")

    # ==================== 核心灵魂：错题深度复盘卡 ====================
    # (a) 上方左右双栏：左列【错误思路】 vs 右列【正确思路】
    # (b) 下方作答核对栏：我的作答 vs 正确答案
    review_table = doc.add_table(rows=2, cols=2)
    review_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_left_indent(review_table, HANG_INDENT_PT)
    review_table.autofit = False

    half_w = OPT_TABLE_WIDTH / 2

    # --- 行 0：上方左右双栏 ---
    # 左单元格：【错误思路】
    c_wrong = review_table.cell(0, 0)
    c_wrong.width = half_w
    set_cell_borders(c_wrong, top="single", bottom="single", left="single", right="single", color="C0C0C0", sz="4")
    set_cell_margins(c_wrong, top=4, bottom=4, left=6, right=6)
    set_cell_shading(c_wrong, color="FAFBFD")

    cp_w = c_wrong.paragraphs[0]
    cp_w.paragraph_format.space_before = Pt(1)
    cp_w.paragraph_format.space_after = Pt(2)
    cp_w.paragraph_format.line_spacing = 1.15
    rw_title = cp_w.add_run("【错误思路】")
    set_font(rw_title, east_asia="黑体", size_pt=10.0, bold=True, color_rgb=RGBColor(0xC0, 0x39, 0x2B))
    rw_hint = cp_w.add_run("（误区诊断/思维盲区）")
    set_font(rw_hint, east_asia="仿宋", size_pt=9.0, color_rgb=RGBColor(0x88, 0x88, 0x88))

    cp_w_body = c_wrong.add_paragraph()
    cp_w_body.paragraph_format.space_before = Pt(1)
    cp_w_body.paragraph_format.space_after = Pt(4)
    cp_w_body.paragraph_format.line_spacing = 1.25
    cp_w_body.add_run("{% if q.wrong_thought %}")
    rw_txt = cp_w_body.add_run("{{ q.wrong_thought }}")
    set_font(rw_txt, east_asia="仿宋", size_pt=9.5, color_rgb=RGBColor(0x44, 0x44, 0x44))
    cp_w_body.add_run("{% else %}")
    rw_empty = cp_w_body.add_run("\n\n\n\n\n")
    set_font(rw_empty, east_asia="仿宋", size_pt=10.0)
    cp_w_body.add_run("{% endif %}")

    # 右单元格：【正确思路】
    c_corr = review_table.cell(0, 1)
    c_corr.width = half_w
    set_cell_borders(c_corr, top="single", bottom="single", left="single", right="single", color="C0C0C0", sz="4")
    set_cell_margins(c_corr, top=4, bottom=4, left=6, right=6)
    set_cell_shading(c_corr, color="FAFBFD")

    cp_c = c_corr.paragraphs[0]
    cp_c.paragraph_format.space_before = Pt(1)
    cp_c.paragraph_format.space_after = Pt(2)
    cp_c.paragraph_format.line_spacing = 1.15
    rc_title = cp_c.add_run("【正确思路】")
    set_font(rc_title, east_asia="黑体", size_pt=10.0, bold=True, color_rgb=RGBColor(0x27, 0xAE, 0x60))
    rc_hint = cp_c.add_run("（破题切入/关键公式）")
    set_font(rc_hint, east_asia="仿宋", size_pt=9.0, color_rgb=RGBColor(0x88, 0x88, 0x88))

    cp_c_body = c_corr.add_paragraph()
    cp_c_body.paragraph_format.space_before = Pt(1)
    cp_c_body.paragraph_format.space_after = Pt(4)
    cp_c_body.paragraph_format.line_spacing = 1.25
    cp_c_body.add_run("{% if q.correct_thought %}")
    rc_txt = cp_c_body.add_run("{{ q.correct_thought }}")
    set_font(rc_txt, east_asia="仿宋", size_pt=9.5, color_rgb=RGBColor(0x44, 0x44, 0x44))
    cp_c_body.add_run("{% else %}")
    rc_empty = cp_c_body.add_run("\n\n\n\n\n")
    set_font(rc_empty, east_asia="仿宋", size_pt=10.0)
    cp_c_body.add_run("{% endif %}")

    # --- 行 1：下方作答核对栏 ---
    c_ans_left = review_table.cell(1, 0)
    c_ans_left.width = half_w
    set_cell_borders(c_ans_left, top="single", bottom="single", left="single", right="none", color="C0C0C0", sz="4")
    set_cell_margins(c_ans_left, top=3, bottom=3, left=6, right=6)
    set_cell_shading(c_ans_left, color="F1F5F9")
    cp_al = c_ans_left.paragraphs[0]
    cp_al.paragraph_format.space_before = Pt(0)
    cp_al.paragraph_format.space_after = Pt(0)
    r_user = cp_al.add_run("我的作答：{{ q.user_answer if q.user_answer else '—' }}")
    set_font(r_user, east_asia="黑体", size_pt=10.0, bold=True, color_rgb=RGBColor(0x2C, 0x3E, 0x50))

    c_ans_right = review_table.cell(1, 1)
    c_ans_right.width = half_w
    set_cell_borders(c_ans_right, top="single", bottom="single", left="none", right="single", color="C0C0C0", sz="4")
    set_cell_margins(c_ans_right, top=3, bottom=3, left=6, right=6)
    set_cell_shading(c_ans_right, color="F1F5F9")
    cp_ar = c_ans_right.paragraphs[0]
    cp_ar.paragraph_format.space_before = Pt(0)
    cp_ar.paragraph_format.space_after = Pt(0)
    cp_ar.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_key = cp_ar.add_run("正确答案：{{ q.answer_text }}")
    set_font(r_key, east_asia="黑体", size_pt=10.0, bold=True, color_rgb=RGBColor(0x1B, 0x4D, 0x3E))

    # ==================== 选项 1：若开启 inline 模式，直接紧随解析 ====================
    doc.add_paragraph("{% if paper.solution_mode == 'inline' %}")
    p_inline_ana = doc.add_paragraph()
    p_inline_ana.paragraph_format.space_before = Pt(4)
    p_inline_ana.paragraph_format.space_after = Pt(6)
    p_inline_ana.paragraph_format.line_spacing = 1.25
    p_inline_ana.paragraph_format.left_indent = Pt(21)
    p_inline_ana.paragraph_format.first_line_indent = Pt(0)
    r_iana_lbl = p_inline_ana.add_run("【官方解析】")
    set_font(r_iana_lbl, east_asia="黑体", size_pt=10.5, bold=True)
    r_iana_txt = p_inline_ana.add_run("{{ q.analysis_text }}")
    set_font(r_iana_txt, east_asia="仿宋", size_pt=10.5, bold=False)

    doc.add_paragraph("{% if q.analysis_images %}{% for aimg in q.analysis_images %}")
    p_iaimg = doc.add_paragraph()
    p_iaimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_iaimg.paragraph_format.space_before = Pt(2)
    p_iaimg.paragraph_format.space_after = Pt(4)
    p_iaimg.paragraph_format.keep_with_next = True
    p_iaimg.add_run("{{ aimg }}")
    doc.add_paragraph("{% endfor %}{% endif %}")
    doc.add_paragraph("{% endif %}")

    # 题目间自然呼吸间距
    p_q_space = doc.add_paragraph()
    p_q_space.paragraph_format.space_before = Pt(2)
    p_q_space.paragraph_format.space_after = Pt(4)

    # 结束题目循环
    doc.add_paragraph("{% endfor %}")

    # ==================== 第 3 节：若为 appendix 模式，附录集中呈现参考答案与深度解析 ====================
    doc.add_paragraph("{% if paper.solution_mode == 'appendix' %}")
    sec3 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    sec3.page_width = PAGE_WIDTH
    sec3.page_height = PAGE_HEIGHT
    sec3.top_margin = Mm(20)
    sec3.bottom_margin = Mm(20)
    sec3.left_margin = MARGIN_LR
    sec3.right_margin = MARGIN_LR

    # 答案速查大标题
    p_ans_head = doc.add_paragraph()
    p_ans_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ans_head.paragraph_format.space_before = Pt(12)
    p_ans_head.paragraph_format.space_after = Pt(12)
    r_ans_head = p_ans_head.add_run("参考答案速查")
    set_font(r_ans_head, east_asia="黑体", size_pt=14, bold=True)

    # 5 列速查答案表格
    doc.add_paragraph("{% for a_row in paper.answer_rows %}")
    ans_table = doc.add_table(rows=1, cols=5)
    ans_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, cell in enumerate(ans_table.rows[0].cells):
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="C0C0C0", sz="4")
        set_cell_shading(cell, color="F8FAFC" if c_idx % 2 == 0 else "FFFFFF")
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(4)
        cr = cp.add_run(f"{{{{ a_row.c{c_idx} }}}}")
        set_font(cr, east_asia="仿宋", size_pt=10.5, bold=False)
    doc.add_paragraph("{% endfor %}")

    # 逐题解析大标题
    p_ana_head = doc.add_paragraph()
    p_ana_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ana_head.paragraph_format.space_before = Pt(20)
    p_ana_head.paragraph_format.space_after = Pt(10)
    r_ana_head = p_ana_head.add_run("参考答案与解析")
    set_font(r_ana_head, east_asia="黑体", size_pt=14, bold=True)

    # 逐题解析循环
    doc.add_paragraph("{% for q in paper.all_questions %}")

    p_item_ans = doc.add_paragraph()
    p_item_ans.paragraph_format.space_before = Pt(7)
    p_item_ans.paragraph_format.space_after = Pt(2)
    p_item_ans.paragraph_format.left_indent = Pt(0)
    p_item_ans.paragraph_format.keep_with_next = True
    r_no = p_item_ans.add_run("{{ q.number }}. 【答案】{{ q.answer_text }}")
    set_font(r_no, east_asia="黑体", size_pt=10.5, bold=True)

    p_item_ana = doc.add_paragraph()
    p_item_ana.paragraph_format.space_before = Pt(1)
    p_item_ana.paragraph_format.space_after = Pt(5)
    p_item_ana.paragraph_format.line_spacing = 1.25
    p_item_ana.paragraph_format.left_indent = Pt(21)
    p_item_ana.paragraph_format.first_line_indent = Pt(0)
    r_ana_lbl = p_item_ana.add_run("【解析】")
    set_font(r_ana_lbl, east_asia="黑体", size_pt=10.5, bold=True)
    r_ana_txt = p_item_ana.add_run("{{ q.analysis_text }}")
    set_font(r_ana_txt, east_asia="仿宋", size_pt=10.5, bold=False)

    doc.add_paragraph("{% if q.analysis_images %}{% for aimg in q.analysis_images %}")
    p_aimg = doc.add_paragraph()
    p_aimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aimg.paragraph_format.space_before = Pt(2)
    p_aimg.paragraph_format.space_after = Pt(4)
    p_aimg.paragraph_format.keep_with_next = True
    p_aimg.add_run("{{ aimg }}")
    doc.add_paragraph("{% endfor %}{% endif %}")

    doc.add_paragraph("{% endfor %}")
    doc.add_paragraph("{% endif %}")

    doc.save(str(docx_path))
    print(f"[✓] 成功创建公考错题本专属模板: {docx_path}")


if __name__ == "__main__":
    create_mistake_template()
