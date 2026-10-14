#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成公考刷题本专属 Word 模板族：
1. template.docx: 完整合订本（封面 + 做题区 + 卷末速查参考答案 + 逐题深度精讲）；
2. template_questions_only.docx: 试卷分册（封面 + 做题区）；
3. template_solutions_only.docx: 答案解析分册（封面 + 卷末速查参考答案 + 逐题深度精讲）。
严格保证三套模板在字号、悬挂缩进、表格尺寸、边框规范上的 100% 统一。
"""

from pathlib import Path
import docx
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls

# 考卷排版核心常量（与国考试卷保持 100% 绝对一致）
HANG_INDENT_PT = 24.0  # 约 2 字符基准悬挂缩进
PAGE_WIDTH = Mm(182)
PAGE_HEIGHT = Mm(257)
MARGIN_LR = Mm(20)
CONTENT_WIDTH = PAGE_WIDTH - MARGIN_LR * 2
OPT_TABLE_WIDTH = CONTENT_WIDTH - Pt(HANG_INDENT_PT)


def set_font(run, east_asia="仿宋", ascii_font="Times New Roman", size_pt=10.5, bold=False, color_rgb=None):
    run.font.name = ascii_font
    if size_pt:
        run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color_rgb:
        run.font.color.rgb = color_rgb
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{ascii_font}" w:eastAsia="{east_asia}" w:hAnsi="{ascii_font}" w:cs="{east_asia}"/>')
        rPr.append(rFonts)
    else:
        rFonts.set(qn("w:eastAsia"), east_asia)
        rFonts.set(qn("w:ascii"), ascii_font)
        rFonts.set(qn("w:hAnsi"), ascii_font)
        rFonts.set(qn("w:cs"), east_asia)


def apply_justified_paragraph(p):
    """为段落设置出版级两端对齐与严格避头尾属性，确保右边缘平齐"""
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="both"/>'))
    pPr.append(parse_xml(f'<w:kinsoku {nsdecls("w")} w:val="1"/>'))
    pPr.append(parse_xml(f'<w:overflowPunct {nsdecls("w")} w:val="0"/>'))


def set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="C5CAD0", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)


def set_table_borders_none(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def set_table_left_indent(table, indent_pt: float):
    tblPr = table._tbl.tblPr
    tblInd = parse_xml(
        f'<w:tblInd {nsdecls("w")} w:w="{int(indent_pt * 20)}" w:type="dxa"/>'
    )
    tblPr.append(tblInd)


def set_cell_margins(cell, top=2, bottom=2, left=0, right=8):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{int(top * 20)}" w:type="dxa"/>\n'
        f'  <w:left w:w="{int(left * 20)}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{int(bottom * 20)}" w:type="dxa"/>\n'
        f'  <w:right w:w="{int(right * 20)}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_cell_shading(cell, color="F8FAFC"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    tcPr.append(shd)


def build_cover_section(doc: docx.Document, title_text: str = "行 测 专 项 刷 题 本", subtitle_suffix: str = ""):
    """构建封面（第 1 节：B5 单栏，标题/信息卡片/名言均匀分布）"""
    sec = doc.sections[0]
    sec.page_width = PAGE_WIDTH
    sec.page_height = PAGE_HEIGHT
    sec.top_margin = Mm(25)
    sec.bottom_margin = Mm(25)
    sec.left_margin = MARGIN_LR
    sec.right_margin = MARGIN_LR

    # 封面大标题
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(40)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run(title_text)
    set_font(r_title, east_asia="黑体", size_pt=26, bold=True)

    # 封面副标题
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(50)
    sub_content = "{{ paper.module_title }}" + (f" · {subtitle_suffix}" if subtitle_suffix else "")
    r_sub = p_sub.add_run(sub_content)
    set_font(r_sub, east_asia="楷体", size_pt=14, bold=False, color_rgb=RGBColor(0x33, 0x33, 0x33))

    # 基本信息卡片表格
    info_table = doc.add_table(rows=3, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    info_widths = [Mm(71), Mm(71)]

    info_data = [
        ("学 员 昵 称：{{ user.nickname }}", "刷 题 题 量：共 {{ paper.question_count }} 题"),
        ("生 成 时 间：{{ user.generated_at }}", "建 议 用 时：{{ paper.suggested_time }} 分钟"),
        ("所 属 题 库：{{ paper.bank_name }}", "考 点 范 围：{{ paper.kp_names }}"),
    ]

    for r_idx, row in enumerate(info_table.rows):
        t0, t1 = info_data[r_idx]
        for c_idx, cell in enumerate(row.cells):
            cell.width = info_widths[c_idx]
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="C5CAD0", sz="5")
            set_cell_shading(cell, color="F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.left_indent = Pt(12)
            text = t0 if c_idx == 0 else t1
            r = p.add_run(text)
            set_font(r, east_asia="仿宋", size_pt=10.5, color_rgb=RGBColor(0x33, 0x33, 0x33))

    # 底部名言
    p_motto = doc.add_paragraph()
    p_motto.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_motto.paragraph_format.space_before = Pt(100)
    p_motto.paragraph_format.space_after = Pt(12)
    r_motto = p_motto.add_run("“ 日拱一卒，功不唐捐；行则将至，做则必成。”")
    set_font(r_motto, east_asia="楷体", size_pt=11.5, color_rgb=RGBColor(0x66, 0x66, 0x66))


def add_page_footer(sec):
    """为指定节添加第 X 页 共 Y 页动态页脚"""
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(6)
    fp.paragraph_format.space_after = Pt(0)
    pPr = fp._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))

    r1 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t xml:space="preserve">第 </w:t></w:r>')
    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
    r2 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t xml:space="preserve"> 页  共 </w:t></w:r>')
    fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>')
    r3 = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="20"/></w:rPr><w:t> 页</w:t></w:r>')
    fp._p.append(r1)
    fp._p.append(fld1)
    fp._p.append(r2)
    fp._p.append(fld2)
    fp._p.append(r3)


def build_questions_section(doc: docx.Document):
    """构建做题正文节（悬挂缩进 + Tab 制表符 + 表格左缩进）"""
    sec2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    sec2.page_width = PAGE_WIDTH
    sec2.page_height = PAGE_HEIGHT
    sec2.top_margin = Mm(20)
    sec2.bottom_margin = Mm(20)
    sec2.left_margin = MARGIN_LR
    sec2.right_margin = MARGIN_LR

    sectPr2 = sec2._sectPr
    cols_xml2 = parse_xml(f'<w:cols {nsdecls("w")} w:num="1"/>')
    sectPr2.append(cols_xml2)
    sectPr2.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>'))

    add_page_footer(sec2)

    # 题目循环开始
    p_q_loop = doc.add_paragraph()
    p_q_loop.add_run("{% for q in paper.all_questions %}")

    # 1. 材料题引导语
    p_mat_intro = doc.add_paragraph()
    p_mat_intro.paragraph_format.space_before = Pt(6)
    p_mat_intro.paragraph_format.space_after = Pt(3)
    p_mat_intro.paragraph_format.left_indent = Pt(HANG_INDENT_PT)
    p_mat_intro.paragraph_format.line_spacing = 1.2
    p_mat_intro.paragraph_format.keep_with_next = True
    r_mat_intro = p_mat_intro.add_run("{% if q.material and q.material.intro %}{{ q.material.intro }}{% endif %}")
    set_font(r_mat_intro, east_asia="黑体", size_pt=10.5, bold=False)

    # 2. 材料题正文
    doc.add_paragraph("{% if q.material and q.material.blocks %}{% for blk in q.material.blocks %}")
    doc.add_paragraph("{% if not blk.is_img %}")
    p_mat_line = doc.add_paragraph()
    apply_justified_paragraph(p_mat_line)
    p_mat_line.paragraph_format.space_before = Pt(1)
    p_mat_line.paragraph_format.space_after = Pt(2)
    p_mat_line.paragraph_format.left_indent = Pt(0)
    p_mat_line.paragraph_format.first_line_indent = Pt(HANG_INDENT_PT)
    p_mat_line.paragraph_format.line_spacing = 1.25
    p_mat_line.paragraph_format.keep_with_next = True
    r_mat_line = p_mat_line.add_run("{{ blk.text }}")
    set_font(r_mat_line, east_asia="仿宋", size_pt=12.0)
    doc.add_paragraph("{% else %}")
    p_mat_img = doc.add_paragraph()
    p_mat_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mat_img.paragraph_format.space_before = Pt(4)
    p_mat_img.paragraph_format.space_after = Pt(4)
    p_mat_img.paragraph_format.keep_with_next = True
    p_mat_img.add_run("{{ blk.img }}")
    doc.add_paragraph("{% endif %}{% endfor %}{% elif q.material and q.material.lines %}{% for line in q.material.lines %}")
    p_mat_fb = doc.add_paragraph()
    apply_justified_paragraph(p_mat_fb)
    p_mat_fb.paragraph_format.space_before = Pt(1)
    p_mat_fb.paragraph_format.space_after = Pt(2)
    p_mat_fb.paragraph_format.left_indent = Pt(0)
    p_mat_fb.paragraph_format.first_line_indent = Pt(HANG_INDENT_PT)
    p_mat_fb.paragraph_format.line_spacing = 1.25
    p_mat_fb.paragraph_format.keep_with_next = True
    r_mat_fb = p_mat_fb.add_run("{{ line }}")
    set_font(r_mat_fb, east_asia="仿宋", size_pt=12.0)
    doc.add_paragraph("{% endfor %}{% endif %}")

    # 3. 题干
    p_stem = doc.add_paragraph()
    apply_justified_paragraph(p_stem)
    p_stem.paragraph_format.space_before = Pt(5)
    p_stem.paragraph_format.space_after = Pt(2)
    p_stem.paragraph_format.line_spacing = 1.25
    p_stem.paragraph_format.keep_with_next = True
    p_stem.paragraph_format.left_indent = Pt(HANG_INDENT_PT)
    p_stem.paragraph_format.first_line_indent = -Pt(HANG_INDENT_PT)
    p_stem.paragraph_format.tab_stops.add_tab_stop(Pt(HANG_INDENT_PT))

    r_no = p_stem.add_run("{{ q.number }}.\t")
    set_font(r_no, east_asia="仿宋", ascii_font="Times New Roman", size_pt=12.0)

    r_tag = p_stem.add_run("{% if q.source_tag %}【{{ q.source_tag }}】{% endif %}")
    set_font(r_tag, east_asia="仿宋", ascii_font="Times New Roman", size_pt=12.0)

    r_stem_0 = p_stem.add_run("{% for seg in q.stem_segments %}{% if seg.is_img %}")
    set_font(r_stem_0, east_asia="仿宋", size_pt=12.0)
    r_stem_img = p_stem.add_run("{{ seg.img }}")
    r_stem_elif = p_stem.add_run("{% elif seg.bold %}")
    set_font(r_stem_elif, east_asia="仿宋", size_pt=12.0)
    r_stem_bold = p_stem.add_run("{{ seg.text }}")
    set_font(r_stem_bold, east_asia="仿宋", size_pt=12.0, bold=True)
    r_stem_else = p_stem.add_run("{% else %}")
    set_font(r_stem_else, east_asia="仿宋", size_pt=12.0)
    r_stem_normal = p_stem.add_run("{{ seg.text }}")
    set_font(r_stem_normal, east_asia="仿宋", size_pt=12.0, bold=False)
    r_stem_end = p_stem.add_run("{% endif %}{% endfor %}")
    set_font(r_stem_end, east_asia="仿宋", size_pt=12.0)

    # 题干配图
    doc.add_paragraph("{% if q.stem_images %}{% for simg in q.stem_images %}")
    p_simg = doc.add_paragraph()
    p_simg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_simg.paragraph_format.space_before = Pt(3)
    p_simg.paragraph_format.space_after = Pt(4)
    p_simg.paragraph_format.keep_with_next = True
    p_simg.add_run("{{ simg }}")
    doc.add_paragraph("{% endfor %}{% endif %}")

    # 4. 4 列选项
    doc.add_paragraph("{% if q.opt_cols == 4 %}")
    t4 = doc.add_table(rows=1, cols=4)
    t4.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders_none(t4)
    set_table_left_indent(t4, HANG_INDENT_PT)
    t4.autofit = False
    col_width_4 = OPT_TABLE_WIDTH / 4
    for idx, c in enumerate(t4.rows[0].cells):
        c.width = col_width_4
        set_cell_margins(c, top=2, bottom=2, left=0, right=6)
        cp = c.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.15
        cp.add_run(f"{{% if q.opt_{idx}_img %}}")
        c_pre = cp.add_run(f"{{{{ q.opt_{idx}_prefix }}}} ")
        set_font(c_pre, east_asia="仿宋", size_pt=12.0)
        cp.add_run(f"{{{{ q.opt_{idx}_img }}}}")
        cp.add_run(f"{{% else %}}")
        cr = cp.add_run(f"{{{{ q.opt_{idx} }}}}")
        set_font(cr, east_asia="仿宋", size_pt=12.0)
        cp.add_run(f"{{% endif %}}")
    doc.add_paragraph("{% endif %}")

    # 5. 2 列选项
    doc.add_paragraph("{% if q.opt_cols == 2 %}")
    t2 = doc.add_table(rows=2, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders_none(t2)
    set_table_left_indent(t2, HANG_INDENT_PT)
    t2.autofit = False
    col_width_2 = OPT_TABLE_WIDTH / 2
    t2_cells = [t2.rows[0].cells[0], t2.rows[0].cells[1], t2.rows[1].cells[0], t2.rows[1].cells[1]]
    for idx, c in enumerate(t2_cells):
        c.width = col_width_2
        set_cell_margins(c, top=2, bottom=2, left=0, right=6)
        cp = c.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.15
        cp.add_run(f"{{% if q.opt_{idx}_img %}}")
        c_pre = cp.add_run(f"{{{{ q.opt_{idx}_prefix }}}} ")
        set_font(c_pre, east_asia="仿宋", size_pt=12.0)
        cp.add_run(f"{{{{ q.opt_{idx}_img }}}}")
        cp.add_run(f"{{% else %}}")
        cr = cp.add_run(f"{{{{ q.opt_{idx} }}}}")
        set_font(cr, east_asia="仿宋", size_pt=12.0)
        cp.add_run(f"{{% endif %}}")
    doc.add_paragraph("{% endif %}")

    # 6. 1 列选项
    doc.add_paragraph("{% if q.opt_cols == 1 %}")
    for i in range(4):
        doc.add_paragraph(f"{{% if q.opt_{i} or q.opt_{i}_img %}}")
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.space_before = Pt(1)
        p_opt.paragraph_format.space_after = Pt(1)
        p_opt.paragraph_format.left_indent = Pt(HANG_INDENT_PT)
        p_opt.paragraph_format.line_spacing = 1.15
        p_opt.add_run(f"{{% if q.opt_{i}_img %}}")
        p_pre = p_opt.add_run(f"{{{{ q.opt_{i}_prefix }}}} ")
        set_font(p_pre, east_asia="仿宋", size_pt=12.0)
        p_opt.add_run(f"{{{{ q.opt_{i}_img }}}}")
        p_opt.add_run(f"{{% else %}}")
        r_opt = p_opt.add_run(f"{{{{ q.opt_{i} }}}}")
        set_font(r_opt, east_asia="仿宋", size_pt=12.0)
        p_opt.add_run(f"{{% endif %}}")
        doc.add_paragraph("{% endif %}")
    doc.add_paragraph("{% endif %}")

    # 结束题目循环
    doc.add_paragraph("{% endfor %}")


def build_solutions_section(doc: docx.Document, is_standalone: bool = False):
    """构建参考答案与解析节"""
    sec3 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    sec3.page_width = PAGE_WIDTH
    sec3.page_height = PAGE_HEIGHT
    sec3.top_margin = Mm(20)
    sec3.bottom_margin = Mm(20)
    sec3.left_margin = MARGIN_LR
    sec3.right_margin = MARGIN_LR

    if is_standalone:
        sectPr3 = sec3._sectPr
        sectPr3.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>'))
    add_page_footer(sec3)

    # 答案速查大标题
    p_ans_head = doc.add_paragraph()
    p_ans_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ans_head.paragraph_format.space_before = Pt(14)
    p_ans_head.paragraph_format.space_after = Pt(10)
    r_ans_head = p_ans_head.add_run("参考答案速查")
    set_font(r_ans_head, east_asia="黑体", size_pt=14, bold=True)

    # 5 列答案速查表
    doc.add_paragraph("{% for a_row in paper.answer_rows %}")
    ans_table = doc.add_table(rows=1, cols=5)
    ans_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, cell in enumerate(ans_table.rows[0].cells):
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="C0C0C0", sz="4")
        set_cell_shading(cell, color="F8FAFC" if c_idx % 2 == 0 else "FFFFFF")
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(4)
        cr = cp.add_run(f"{{{{ a_row.c{c_idx} }}}}")
        set_font(cr, east_asia="仿宋", size_pt=10.5, bold=False)
    doc.add_paragraph("{% endfor %}")

    # 逐题解析大标题
    p_ana_head = doc.add_paragraph()
    p_ana_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ana_head.paragraph_format.space_before = Pt(20)
    p_ana_head.paragraph_format.space_after = Pt(10)
    r_ana_head = p_ana_head.add_run("参考答案与深度精解")
    set_font(r_ana_head, east_asia="黑体", size_pt=14, bold=True)

    # 逐题解析循环
    doc.add_paragraph("{% for q in paper.all_questions %}")

    p_item_ans = doc.add_paragraph()
    p_item_ans.paragraph_format.space_before = Pt(7)
    p_item_ans.paragraph_format.space_after = Pt(2)
    p_item_ans.paragraph_format.left_indent = Pt(0)
    p_item_ans.paragraph_format.keep_with_next = True
    r_no = p_item_ans.add_run("{{ q.number }}. 【答案】{{ q.answer_text }}")
    set_font(r_no, east_asia="黑体", size_pt=10.5, bold=True)

    p_item_ana = doc.add_paragraph()
    apply_justified_paragraph(p_item_ana)
    p_item_ana.paragraph_format.space_before = Pt(1)
    p_item_ana.paragraph_format.space_after = Pt(5)
    p_item_ana.paragraph_format.line_spacing = 1.25
    p_item_ana.paragraph_format.left_indent = Pt(21)
    p_item_ana.paragraph_format.first_line_indent = Pt(0)
    r_ana_lbl = p_item_ana.add_run("【解析】")
    set_font(r_ana_lbl, east_asia="黑体", size_pt=10.5, bold=True)

    r_ana_seg_start = p_item_ana.add_run("{% if q.analysis_segments %}{% for seg in q.analysis_segments %}{% if seg.is_img %}")
    set_font(r_ana_seg_start, east_asia="仿宋", size_pt=10.5)
    r_ana_img = p_item_ana.add_run("{{ seg.img }}")
    r_ana_elif = p_item_ana.add_run("{% elif seg.bold %}")
    set_font(r_ana_elif, east_asia="仿宋", size_pt=10.5)
    r_ana_bold = p_item_ana.add_run("{{ seg.text }}")
    set_font(r_ana_bold, east_asia="仿宋", size_pt=10.5, bold=True)
    r_ana_else = p_item_ana.add_run("{% else %}")
    set_font(r_ana_else, east_asia="仿宋", size_pt=10.5)
    r_ana_normal = p_item_ana.add_run("{{ seg.text }}")
    set_font(r_ana_normal, east_asia="仿宋", size_pt=10.5, bold=False)
    r_ana_seg_end = p_item_ana.add_run("{% endif %}{% endfor %}{% else %}")
    set_font(r_ana_seg_end, east_asia="仿宋", size_pt=10.5)
    r_ana_txt = p_item_ana.add_run("{{ q.analysis_text }}")
    set_font(r_ana_txt, east_asia="仿宋", size_pt=10.5, bold=False)
    r_ana_close = p_item_ana.add_run("{% endif %}")
    set_font(r_ana_close, east_asia="仿宋", size_pt=10.5)

    # 解析配图
    doc.add_paragraph("{% if q.analysis_images %}{% for aimg in q.analysis_images %}")
    p_aimg = doc.add_paragraph()
    p_aimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aimg.paragraph_format.space_before = Pt(2)
    p_aimg.paragraph_format.space_after = Pt(4)
    p_aimg.paragraph_format.keep_with_next = True
    p_aimg.add_run("{{ aimg }}")
    doc.add_paragraph("{% endfor %}{% endif %}")

    doc.add_paragraph("{% endfor %}")


def create_all_templates():
    target_dir = Path(r"backend/plugin/render_book/templates/gongkao_practice/1.0.0")
    target_dir.mkdir(parents=True, exist_ok=True)

    # 1. 完整合订版 (template.docx)
    doc_merged = docx.Document()
    build_cover_section(doc_merged, title_text="行 测 专 项 刷 题 本")
    build_questions_section(doc_merged)
    build_solutions_section(doc_merged, is_standalone=False)
    p_merged = target_dir / "template.docx"
    doc_merged.save(str(p_merged))
    print(f"[✓] 成功创建公考刷题本·合订版模板: {p_merged}")

    # 2. 试卷分册 (template_questions_only.docx)
    doc_q = docx.Document()
    build_cover_section(doc_q, title_text="行 测 专 项 刷 题 本", subtitle_suffix="试题分册")
    build_questions_section(doc_q)
    p_q = target_dir / "template_questions_only.docx"
    doc_q.save(str(p_q))
    print(f"[✓] 成功创建公考刷题本·试卷分册模板: {p_q}")

    # 3. 答案解析分册 (template_solutions_only.docx)
    doc_s = docx.Document()
    build_cover_section(doc_s, title_text="行 测 专 项 刷 题 本", subtitle_suffix="答案解析分册")
    build_solutions_section(doc_s, is_standalone=True)
    p_s = target_dir / "template_solutions_only.docx"
    doc_s.save(str(p_s))
    print(f"[✓] 成功创建公考刷题本·答案解析分册模板: {p_s}")


if __name__ == "__main__":
    create_all_templates()
